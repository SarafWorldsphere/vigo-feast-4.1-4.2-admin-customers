import asyncio
import json
from io import BytesIO
from datetime import datetime, timezone
from datetime import date, timedelta
from pathlib import Path
from typing import Any, Literal
from uuid import UUID, uuid4

from fastapi import Depends, FastAPI, Header, HTTPException, Query, Request
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, ConfigDict, Field
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb

from app.config import settings
from app.auth import (
    create_admin_session,
    create_admin_token,
    decode_token,
    find_admin_by_mobile,
    find_admin_by_username,
    generate_otp,
    save_otp_code,
    send_otp_message,
    verify_otp_code,
    verify_password,
)
from app.db import check_db_connection, get_db_connection
from app.order_status import CANCELLED_ORDER_SQL

app = FastAPI(title="Vigo Feast Admin Console")

PROJECT_ROOT = Path(__file__).resolve().parent.parent
app.mount(
    "/static",
    StaticFiles(directory=PROJECT_ROOT / "static"),
    name="static",
)


@app.get("/")
async def admin_console() -> FileResponse:
    return FileResponse(PROJECT_ROOT / "templates" / "admin_console.html")


class AdminPasswordLoginRequest(BaseModel):
    username: str = Field(..., min_length=1)
    password: str = Field(..., min_length=1)


class AdminOtpRequest(BaseModel):
    mobile_number: str = Field(..., min_length=1)
    language: Literal["english", "hindi", "telugu"] = "english"


class AdminOtpVerifyRequest(BaseModel):
    mobile_number: str = Field(..., min_length=1)
    otp_code: str = Field(..., min_length=4, max_length=6)


class CustomerUpdateRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    name: str | None = Field(default=None, min_length=1)
    mobile: str | None = None
    email: str | None = None
    default_address_id: int | None = Field(default=None, gt=0)


class CustomerStatusRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    is_active: bool


def require_admin_auth(
    request: Request,
    authorization: str | None = Header(default=None),
) -> dict[str, Any]:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Missing bearer token")

    token = authorization.split(" ", 1)[1].strip()
    try:
        payload = decode_token(token)
    except ValueError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc

    role = payload.get("role")
    if not role:
        raise HTTPException(status_code=401, detail="Token missing role")

    action_column = {
        "GET": "can_view",
        "HEAD": "can_view",
        "POST": "can_create",
        "PUT": "can_update",
        "PATCH": "can_update",
        "DELETE": "can_delete",
    }.get(request.method)
    if action_column is None:
        raise HTTPException(status_code=403, detail="This action is not permitted")
    permission_found = False
    permission_granted = False
    franchisee_id = payload.get("franchisee_id")
    if isinstance(franchisee_id, str):
        try:
            franchisee_id = UUID(franchisee_id)
        except ValueError:
            pass
    franchisee_filter = "" if isinstance(franchisee_id, UUID) else " AND franchisee_id = %s"
    permission_parameters = (
        (role,) if isinstance(franchisee_id, UUID) else (role, franchisee_id)
    )
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT can_view, can_create, can_update, can_delete
                FROM vgf_adm_role_permission
                WHERE role_name = %s
                  AND permission_name = 'admin_console'
                  {franchisee_filter}
                  AND is_active IS TRUE
                """,
                permission_parameters,
            )
            permission = cur.fetchone()
            if permission:
                permission_found = True
                permission_index = {
                    "can_view": 0,
                    "can_create": 1,
                    "can_update": 2,
                    "can_delete": 3,
                }[action_column]
                permission_granted = bool(permission[permission_index])
    if permission_found and not permission_granted:
        raise HTTPException(status_code=403, detail="This action is not permitted")
    if not permission_found and str(role).lower() != "admin":
        raise HTTPException(status_code=403, detail="This role has no admin permissions")

    return payload


def _count_active_customers() -> int:
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT COUNT(*) FROM vgf_customers WHERE is_active = TRUE AND is_deleted = FALSE"
            )
            row = cur.fetchone()
            return int(row[0] or 0)


def _count_live_orders() -> int:
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT COUNT(*) FROM vgf_orders WHERE is_active = TRUE AND is_deleted = FALSE AND order_status NOT IN ('delivered', 'cancelled', 'canceled')"
            )
            row = cur.fetchone()
            return int(row[0] or 0)


def _sum_today_revenue() -> float:
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT COALESCE(SUM(total_amount), 0) FROM vgf_orders WHERE is_active = TRUE AND is_deleted = FALSE AND created_at >= CURRENT_DATE"
            )
            row = cur.fetchone()
            return float(row[0] or 0)


def _count_active_sessions() -> int:
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT COUNT(*) FROM vgf_adm_sessions WHERE is_active = TRUE AND is_deleted = FALSE AND expires_at > NOW()"
            )
            row = cur.fetchone()
            return int(row[0] or 0)


def _percent_change(current: float, previous: float) -> float | None:
    if previous == 0:
        return 0 if current == 0 else None
    return round((current - previous) * 100 / previous, 1)


def _dashboard_trends() -> dict[str, float | None]:
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT
                    (SELECT COUNT(*)
                     FROM vgf_customers
                     WHERE is_active IS TRUE AND is_deleted IS FALSE
                       AND created_at >= date_trunc('week', NOW())) AS customers_this_week,
                    (SELECT COUNT(*)
                     FROM vgf_customers
                     WHERE is_active IS TRUE AND is_deleted IS FALSE
                       AND created_at >= date_trunc('week', NOW()) - INTERVAL '7 days'
                       AND created_at < date_trunc('week', NOW())) AS customers_last_week,
                    (SELECT COUNT(*)
                     FROM vgf_orders
                     WHERE is_active IS TRUE AND is_deleted IS FALSE
                       AND order_status NOT IN ('delivered', 'cancelled', 'canceled')
                       AND created_at >= NOW() - INTERVAL '1 hour'
                       AND created_at < NOW()) AS live_orders_this_hour,
                    (SELECT COUNT(*)
                     FROM vgf_orders
                     WHERE is_active IS TRUE AND is_deleted IS FALSE
                       AND order_status NOT IN ('delivered', 'cancelled', 'canceled')
                       AND created_at >= NOW() - INTERVAL '2 hours'
                       AND created_at < NOW() - INTERVAL '1 hour') AS live_orders_last_hour,
                    (SELECT COALESCE(SUM(total_amount), 0)
                     FROM vgf_orders
                     WHERE is_active IS TRUE AND is_deleted IS FALSE
                       AND created_at >= CURRENT_DATE AND created_at < NOW()) AS revenue_today,
                    (SELECT COALESCE(SUM(total_amount), 0)
                     FROM vgf_orders
                     WHERE is_active IS TRUE AND is_deleted IS FALSE
                       AND created_at >= CURRENT_DATE - INTERVAL '1 day'
                       AND created_at < NOW() - INTERVAL '1 day') AS revenue_yesterday,
                    (SELECT COUNT(*)
                     FROM vgf_adm_sessions s
                     WHERE s.is_active IS TRUE AND s.is_deleted IS FALSE
                       AND s.expires_at > NOW()) AS sessions_now,
                    (SELECT COUNT(*)
                     FROM vgf_adm_sessions s
                     WHERE s.is_deleted IS FALSE
                       AND s.created_at <= NOW() - INTERVAL '1 hour'
                       AND s.expires_at > NOW() - INTERVAL '1 hour'
                       AND (
                           s.is_active IS TRUE
                           OR EXISTS (
                               SELECT 1 FROM vgf_adm_logs l
                               WHERE l.entity_type = 'admin_session'
                                 AND l.action = 'admin.logout'
                                 AND l.entity_id = s.id
                                 AND l.created_at >= NOW() - INTERVAL '1 hour'
                           )
                       )) AS sessions_one_hour_ago
                """
            )
            row = cur.fetchone()

    return {
        "active_customers": _percent_change(float(row[0]), float(row[1])),
        "live_orders": _percent_change(float(row[2]), float(row[3])),
        "todays_revenue": _percent_change(float(row[4]), float(row[5])),
        "active_admin_sessions": _percent_change(float(row[6]), float(row[7])),
    }


def _dashboard_supporting_data() -> dict[str, Any]:
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                """
                SELECT days.day::date AS order_date,
                       COUNT(DISTINCT o.order_id) AS order_count,
                       COALESCE(SUM(p.amount), 0) AS revenue,
                       days.day::date = CURRENT_DATE AS is_today
                FROM generate_series(
                    CURRENT_DATE - INTERVAL '6 days',
                    CURRENT_DATE,
                    INTERVAL '1 day'
                ) AS days(day)
                LEFT JOIN vgf_orders o
                  ON o.created_at::date = days.day::date
                 AND o.is_active IS TRUE AND o.is_deleted IS FALSE
                LEFT JOIN vgf_payments p
                  ON p.order_id = o.order_id
                 AND p.payment_status = 'completed'
                 AND p.is_active IS TRUE AND p.is_deleted IS FALSE
                GROUP BY days.day
                ORDER BY days.day
                """
            )
            orders_last_7_days = [
                {
                    "date": row["order_date"].isoformat(),
                    "order_count": int(row["order_count"]),
                    "revenue": float(row["revenue"]),
                    "is_today": row["is_today"],
                }
                for row in cur.fetchall()
            ]
            cur.execute(
                """
                SELECT l.action, l.entity_type, l.entity_name, l.description,
                       l.created_at,
                       COALESCE(
                           u.full_name,
                           (
                               SELECT a.full_name
                               FROM vgf_admin a
                               WHERE a.username = l.added_by
                                  OR a.email = l.added_by
                                  OR a.mobile_no = l.added_by
                               LIMIT 1
                           ),
                           l.added_by
                       ) AS admin_name
                FROM vgf_adm_logs l
                LEFT JOIN vgf_users u ON u.user_id = l.user_id
                WHERE l.is_deleted IS FALSE
                ORDER BY l.created_at DESC
                LIMIT 10
                """
            )
            recent_activity = [
                {
                    "action": row["action"],
                    "entity_type": row["entity_type"],
                    "entity_name": row["entity_name"],
                    "description": row["description"],
                    "created_at": row["created_at"].isoformat()
                    if row["created_at"]
                    else None,
                    "admin_name": row["admin_name"],
                }
                for row in cur.fetchall()
            ]
    return {
        "orders_last_7_days": orders_last_7_days,
        "recent_activity": recent_activity,
    }


def _build_dashboard_payload() -> dict[str, Any]:
    active_customers = _count_active_customers()
    live_orders = _count_live_orders()
    todays_revenue = _sum_today_revenue()
    active_sessions = _count_active_sessions()
    trends = _dashboard_trends()

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        **_dashboard_supporting_data(),
        "kpis": [
            {
                "key": "active_customers",
                "title": "Active Customers",
                "value": active_customers,
                "change_vs_last_week": trends["active_customers"],
                "unit": "count",
                "source_table": "vgf_customers",
            },
            {
                "key": "live_orders",
                "title": "Live Orders",
                "value": live_orders,
                "change_in_last_hour": trends["live_orders"],
                "unit": "count",
                "source_table": "vgf_orders",
            },
            {
                "key": "todays_revenue",
                "title": "Today's Revenue",
                "value": todays_revenue,
                "change_vs_yesterday": trends["todays_revenue"],
                "currency": "INR",
                "unit": "amount",
                "source_table": "vgf_orders",
            },
            {
                "key": "active_admin_sessions",
                "title": "Active Admin Sessions",
                "value": active_sessions,
                "change_vs_last_hour": trends["active_admin_sessions"],
                "unit": "count",
                "source_table": "vgf_adm_sessions",
            },
        ],
    }


def _kpi_detail_active_customers() -> dict[str, Any]:
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT customer_id, name, email, mobile, loyalty_points FROM vgf_customers WHERE is_active = TRUE AND is_deleted = FALSE ORDER BY customer_id DESC"
            )
            rows = cur.fetchall()
    return {
        "key": "active_customers",
        "title": "Active Customers",
        "source_table": "vgf_customers",
        "count": len(rows),
        "data": [
            {
                "customer_id": row[0],
                "name": row[1],
                "email": row[2],
                "mobile": row[3],
                "loyalty_points": row[4],
            }
            for row in rows
        ],
    }


def _kpi_detail_live_orders(order_id: int | None = None) -> dict[str, Any]:
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                """
                SELECT o.order_status AS status, COUNT(*) AS order_count
                FROM vgf_orders o
                WHERE o.is_active IS TRUE AND o.is_deleted IS FALSE
                  AND o.order_status NOT IN ('delivered', 'cancelled', 'canceled')
                GROUP BY o.order_status
                ORDER BY o.order_status
                """
            )
            status_breakdown = [
                {"status": row["status"], "count": int(row["order_count"])}
                for row in cur.fetchall()
            ]
            cur.execute(
                """
                SELECT hours.hour AS order_hour, COUNT(o.order_id) AS order_count
                FROM generate_series(
                    date_trunc('hour', NOW()) - INTERVAL '23 hours',
                    date_trunc('hour', NOW()),
                    INTERVAL '1 hour'
                ) AS hours(hour)
                LEFT JOIN vgf_orders o
                  ON date_trunc('hour', o.created_at) = hours.hour
                 AND o.created_at < NOW()
                 AND o.is_active IS TRUE AND o.is_deleted IS FALSE
                 AND o.order_status NOT IN ('delivered', 'cancelled', 'canceled')
                GROUP BY hours.hour
                ORDER BY hours.hour
                """
            )
            orders_per_hour = [
                {
                    "hour": row["order_hour"].isoformat(),
                    "order_count": int(row["order_count"]),
                }
                for row in cur.fetchall()
            ]
            order_query = """
                SELECT o.order_id, c.name AS customer_name, o.restaurant_id,
                       o.total_amount, o.order_status, o.delivery_eta
                FROM vgf_orders o
                LEFT JOIN vgf_customers c ON c.customer_id = o.customer_id
                WHERE o.is_active IS TRUE AND o.is_deleted IS FALSE
                  AND o.order_status NOT IN ('delivered', 'cancelled', 'canceled')
            """
            order_parameters: tuple[Any, ...] = ()
            if order_id is not None:
                order_query += " ORDER BY (o.order_id = %s) DESC"
                order_parameters = (order_id,)
            else:
                order_query += " ORDER BY o.created_at DESC, o.order_id DESC"
            if order_id is not None:
                order_query += ", o.created_at DESC, o.order_id DESC"
            cur.execute(order_query, order_parameters)
            rows = cur.fetchall()
    return {
        "key": "live_orders",
        "title": "Live Orders",
        "source_table": "vgf_orders",
        "visualization": "live_orders",
        "count": sum(item["count"] for item in status_breakdown),
        "status_breakdown": status_breakdown,
        "orders_per_hour": orders_per_hour,
        "data": [{
            "order_id": row["order_id"],
            "customer_name": row["customer_name"],
            "restaurant_id": str(row["restaurant_id"]) if row["restaurant_id"] else None,
            "amount": float(row["total_amount"] or 0),
            "order_status": row["order_status"],
            "delivery_eta": row["delivery_eta"].isoformat()
            if row["delivery_eta"]
            else None,
        } for row in rows],
    }


def _kpi_detail_todays_revenue() -> dict[str, Any]:
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                """
                SELECT hours.hour AS revenue_hour,
                       COALESCE(SUM(o.total_amount), 0) AS revenue
                FROM generate_series(
                    date_trunc('day', NOW()),
                    date_trunc('hour', NOW()),
                    INTERVAL '1 hour'
                ) AS hours(hour)
                LEFT JOIN vgf_orders o
                  ON date_trunc('hour', o.created_at) = hours.hour
                 AND o.created_at < NOW()
                 AND o.is_active IS TRUE AND o.is_deleted IS FALSE
                GROUP BY hours.hour
                ORDER BY hours.hour
                """
            )
            hourly_revenue = [
                {
                    "hour": row["revenue_hour"].isoformat(),
                    "revenue": float(row["revenue"] or 0),
                }
                for row in cur.fetchall()
            ]
            cur.execute(
                """
                SELECT payment_id, order_id, method, amount, transaction_ref,
                       payment_status, gateway_name, created_at
                FROM vgf_payments
                WHERE created_at >= CURRENT_DATE
                  AND created_at < NOW()
                  AND is_active IS TRUE AND is_deleted IS FALSE
                ORDER BY created_at DESC, payment_id DESC
                """
            )
            payments = [
                {
                    "payment_id": row["payment_id"],
                    "order_id": row["order_id"],
                    "method": row["method"],
                    "amount": float(row["amount"] or 0),
                    "transaction_ref": row["transaction_ref"],
                    "payment_status": row["payment_status"],
                    "gateway_name": row["gateway_name"],
                    "created_at": row["created_at"].isoformat()
                    if row["created_at"]
                    else None,
                }
                for row in cur.fetchall()
            ]
    total = sum(float(point["revenue"]) for point in hourly_revenue)
    return {
        "key": "todays_revenue",
        "title": "Today's Revenue",
        "source_table": "vgf_orders",
        "visualization": "todays_revenue",
        "total": total,
        "hourly_revenue": hourly_revenue,
        "data": payments,
    }


def _kpi_detail_active_admin_sessions() -> dict[str, Any]:
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT COALESCE(a.full_name, u.full_name, s.added_by),
                       s.ip_address, s.expires_at
                FROM vgf_adm_sessions s
                LEFT JOIN vgf_users u ON u.user_id = s.user_id
                LEFT JOIN vgf_admin a ON a.email = u.email OR a.mobile_no = u.mobile
                WHERE s.is_active IS TRUE AND s.is_deleted IS FALSE
                  AND s.expires_at > NOW()
                ORDER BY s.created_at DESC
                """
            )
            rows = cur.fetchall()
    return {
        "key": "active_admin_sessions",
        "title": "Active Admin Sessions",
        "source_table": "vgf_adm_sessions",
        "count": len(rows),
        "data": [
            {
                "full_name": row[0],
                "ip_address": row[1],
                "expires_at": row[2].isoformat() if row[2] else None,
            }
            for row in rows
        ],
    }


def _get_kpi_detail(card_key: str) -> dict[str, Any]:
    lookup = {
        "active_customers": _kpi_detail_active_customers,
        "live_orders": _kpi_detail_live_orders,
        "todays_revenue": _kpi_detail_todays_revenue,
        "active_admin_sessions": _kpi_detail_active_admin_sessions,
    }
    if card_key not in lookup:
        raise HTTPException(status_code=404, detail="KPI card not found")
    return lookup[card_key]()


_CUSTOMER_SELECT = """
    SELECT customer_id, user_id, name, mobile, email, default_address_id,
           loyalty_points, created_at, updated_at, franchisee_id,
           is_active, is_deleted
    FROM vgf_customers
"""


def _customer_snapshot(customer: dict[str, Any]) -> dict[str, Any]:
    return json.loads(json.dumps(customer, default=str))


def _fetch_customer(
    cur: Any, customer_id: int, *, for_update: bool = False
) -> dict[str, Any] | None:
    query = _CUSTOMER_SELECT + """
        WHERE customer_id = %s AND is_deleted = FALSE
    """
    if for_update:
        query += " FOR UPDATE"
    cur.execute(query, (customer_id,))
    return cur.fetchone()


def _write_customer_audit(
    cur: Any,
    *,
    customer: dict[str, Any],
    current_admin: dict[str, Any],
    request: Request,
    action: str,
    old_data: dict[str, Any],
    new_data: dict[str, Any],
) -> None:
    admin_subject = str(current_admin.get("sub") or "admin")
    changed_by = None
    admin_id = current_admin.get("admin_id")
    if isinstance(admin_id, int):
        cur.execute(
            """
            SELECT u.user_id
            FROM vgf_admin a
            JOIN vgf_users u ON u.email = a.email OR u.mobile = a.mobile_no
            WHERE a.admin_id = %s
            LIMIT 1
            """,
            (admin_id,),
        )
        user = cur.fetchone()
        changed_by = user["user_id"] if user else None

    franchisee_id = customer["franchisee_id"]
    if franchisee_id is not None:
        cur.execute("SELECT 1 FROM vgf_franchisees WHERE id = %s", (franchisee_id,))
        if not cur.fetchone():
            franchisee_id = None

    now = datetime.now(timezone.utc)
    client_ip = request.client.host if request.client else None
    cur.execute(
        """
        INSERT INTO vgf_customer_audit (
            franchisee_id, customer_id, action, old_data, new_data, changed_by,
            added_by, added_date, added_time, added_ip,
            last_updated_by, last_updated_date, last_updated_time, last_updated_ip
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            franchisee_id,
            customer["customer_id"],
            action,
            Jsonb(_customer_snapshot(old_data)),
            Jsonb(_customer_snapshot(new_data)),
            changed_by,
            admin_subject,
            now.date(),
            now.time().replace(microsecond=0),
            client_ip,
            admin_subject,
            now.date(),
            now.time().replace(microsecond=0),
            client_ip,
        ),
    )


def _write_admin_log(
    cur: Any,
    *,
    current_admin: dict[str, Any],
    action: str,
    entity_type: str,
    description: str,
    ip_address: str | None,
    entity_id: UUID | None = None,
    entity_name: str | None = None,
    old_data: dict[str, Any] | None = None,
    new_data: dict[str, Any] | None = None,
) -> None:
    session_id = UUID(str(current_admin["session_id"]))
    cur.execute(
        "SELECT user_id, franchisee_id FROM vgf_adm_sessions WHERE id = %s",
        (session_id,),
    )
    session = cur.fetchone()
    if not session:
        raise HTTPException(status_code=401, detail="Admin session not found")
    session_user_id = (
        session["user_id"] if isinstance(session, dict) else session[0]
    )
    session_franchisee_id = (
        session["franchisee_id"] if isinstance(session, dict) else session[1]
    )

    subject = str(current_admin.get("sub") or "admin")
    cur.execute(
        """
        INSERT INTO vgf_adm_logs (
            id, franchisee_id, user_id, action, entity_type, entity_id,
            description, ip_address, created_at, added_by, added_date,
            added_time, added_ip, last_updated_by, last_updated_date,
            last_updated_time, last_updated_ip, is_active, is_deleted,
            entity_name, old_data, new_data
        )
        VALUES (
            %s, %s, %s, %s, %s, %s, %s, %s, NOW(), %s, CURRENT_DATE,
            CURRENT_TIME, %s, %s, CURRENT_DATE, CURRENT_TIME, %s, TRUE,
            FALSE, %s, %s, %s
        )
        """,
        (
            uuid4(),
            session_franchisee_id,
            session_user_id,
            action,
            entity_type,
            entity_id,
            description,
            ip_address,
            subject,
            ip_address,
            subject,
            ip_address,
            entity_name,
            Jsonb(old_data) if old_data is not None else None,
            Jsonb(new_data) if new_data is not None else None,
        ),
    )


def _customer_list_filters(
    search: str | None,
    status: Literal["all", "active", "inactive"],
    joined_from: date | None,
    joined_to: date | None,
) -> tuple[str, list[Any]]:
    if joined_from and joined_to and joined_from > joined_to:
        raise HTTPException(
            status_code=422,
            detail="joined_from must be on or before joined_to",
        )
    conditions = ["c.is_deleted = FALSE"]
    parameters: list[Any] = []
    if status == "active":
        conditions.append("c.is_active IS TRUE")
    elif status == "inactive":
        conditions.append("c.is_active IS NOT TRUE")
    if search and search.strip():
        pattern = f"%{search.strip()}%"
        conditions.append("(c.name ILIKE %s OR c.mobile ILIKE %s OR c.email ILIKE %s)")
        parameters.extend((pattern, pattern, pattern))
    if joined_from:
        conditions.append("c.created_at >= %s")
        parameters.append(joined_from)
    if joined_to:
        conditions.append("c.created_at < %s")
        parameters.append(joined_to + timedelta(days=1))
    return " AND ".join(conditions), parameters


@app.get("/admin/customers")
async def list_customers(
    current_admin: dict[str, Any] = Depends(require_admin_auth),
    search: str | None = Query(default=None, max_length=200),
    status: Literal["all", "active", "inactive"] = "all",
    joined_from: date | None = None,
    joined_to: date | None = None,
    focus_customer_id: int | None = Query(default=None, ge=1),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=25, ge=1, le=100),
):
    del current_admin
    where_clause, parameters = _customer_list_filters(
        search, status, joined_from, joined_to
    )
    effective_page = page
    if focus_customer_id is not None:
        where_clause = "c.is_deleted = FALSE"
        parameters = []

    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            if focus_customer_id is not None:
                cur.execute(
                    """
                    SELECT created_at
                    FROM vgf_customers
                    WHERE customer_id = %s AND is_deleted IS FALSE
                    """,
                    (focus_customer_id,),
                )
                focus_customer = cur.fetchone()
                if focus_customer:
                    cur.execute(
                        """
                        SELECT COUNT(*) AS customers_before
                        FROM vgf_customers
                        WHERE is_deleted IS FALSE
                          AND (
                              created_at > %s
                              OR (created_at = %s AND customer_id > %s)
                          )
                        """,
                        (
                            focus_customer["created_at"],
                            focus_customer["created_at"],
                            focus_customer_id,
                        ),
                    )
                    effective_page = (
                        cur.fetchone()["customers_before"] // page_size
                    ) + 1
            offset = (effective_page - 1) * page_size
            cur.execute(
                f"SELECT COUNT(*) AS total FROM vgf_customers c WHERE {where_clause}",
                parameters,
            )
            total = cur.fetchone()["total"]
            cur.execute(
                f"""
                SELECT c.customer_id, c.name, c.mobile, c.email,
                       c.loyalty_points, c.is_active, c.created_at,
                       c.default_address_id,
                       COALESCE(o.order_count, 0) AS order_count,
                       COALESCE(o.cancelled_order_count, 0) AS cancelled_order_count
                FROM vgf_customers c
                LEFT JOIN (
                    SELECT o.customer_id,
                           COUNT(*) FILTER (
                               WHERE NOT ({CANCELLED_ORDER_SQL})
                           ) AS order_count,
                           COUNT(*) FILTER (
                               WHERE ({CANCELLED_ORDER_SQL})
                           ) AS cancelled_order_count
                    FROM vgf_orders o
                    WHERE o.is_active IS TRUE AND o.is_deleted IS FALSE
                    GROUP BY o.customer_id
                ) o ON o.customer_id = c.customer_id
                WHERE {where_clause}
                ORDER BY c.created_at DESC, c.customer_id DESC
                LIMIT %s OFFSET %s
                """,
                [*parameters, page_size, offset],
            )
            customers = cur.fetchall()

    return {
        "items": customers,
        "page": effective_page,
        "page_size": page_size,
        "total": total,
        "total_pages": (total + page_size - 1) // page_size,
    }


@app.get("/admin/customers/export.xlsx")
async def export_customers(
    current_admin: dict[str, Any] = Depends(require_admin_auth),
    search: str | None = Query(default=None, max_length=200),
    status: Literal["all", "active", "inactive"] = "all",
    joined_from: date | None = None,
    joined_to: date | None = None,
):
    del current_admin
    where_clause, parameters = _customer_list_filters(
        search, status, joined_from, joined_to
    )
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                f"""
                SELECT c.customer_id, c.name, c.mobile, c.email,
                       c.loyalty_points, c.is_active, c.created_at,
                       COALESCE(o.order_count, 0) AS order_count,
                       COALESCE(o.cancelled_order_count, 0) AS cancelled_order_count
                FROM vgf_customers c
                LEFT JOIN (
                    SELECT o.customer_id,
                           COUNT(*) FILTER (
                               WHERE NOT ({CANCELLED_ORDER_SQL})
                           ) AS order_count,
                           COUNT(*) FILTER (
                               WHERE ({CANCELLED_ORDER_SQL})
                           ) AS cancelled_order_count
                    FROM vgf_orders o
                    WHERE o.is_active IS TRUE AND o.is_deleted IS FALSE
                    GROUP BY o.customer_id
                ) o ON o.customer_id = c.customer_id
                WHERE {where_clause}
                ORDER BY c.created_at DESC, c.customer_id DESC
                """,
                parameters,
            )
            customers = cur.fetchall()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Customers"
    worksheet.append((
        "Customer ID", "Name", "Mobile", "Email", "Orders", "Cancelled",
        "Loyalty Points", "Status", "Joined",
    ))
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    worksheet.freeze_panes = "A2"
    for row in customers:
        worksheet.append((
            row["customer_id"],
            row["name"],
            str(row["mobile"]) if row["mobile"] is not None else "",
            row["email"] or "",
            row["order_count"],
            row["cancelled_order_count"],
            row["loyalty_points"],
            "Active" if row["is_active"] else "Inactive",
            row["created_at"].replace(tzinfo=None)
            if row["created_at"] is not None
            else None,
        ))
        worksheet.cell(worksheet.max_row, 3).number_format = "@"
        worksheet.cell(worksheet.max_row, 9).number_format = "dd/mm/yyyy hh:mm"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column, width in {
        "A": 14, "B": 24, "C": 18, "D": 34, "E": 12,
        "F": 12, "G": 18, "H": 12, "I": 22,
    }.items():
        worksheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": 'attachment; filename="customers.xlsx"'},
    )


@app.get("/admin/customers/summary")
async def customer_summary(
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    del current_admin
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                f"""
                WITH customer_counts AS (
                    SELECT
                        COUNT(*) AS total,
                        COUNT(*) FILTER (WHERE is_active IS NOT TRUE) AS inactive,
                        COUNT(*) FILTER (
                            WHERE created_at >= date_trunc('week', NOW())
                        ) AS new_this_week,
                        COUNT(*) FILTER (
                            WHERE created_at >= date_trunc('week', NOW()) - INTERVAL '7 days'
                              AND created_at < date_trunc('week', NOW())
                        ) AS new_last_week
                    FROM vgf_customers
                    WHERE is_deleted IS FALSE
                ),
                ratings AS (
                    SELECT AVG(rating)::float AS average_rating, COUNT(*) AS review_count
                    FROM vgf_customer_feedback
                    WHERE is_active IS TRUE AND is_deleted IS FALSE
                ),
                order_counts AS (
                    SELECT
                        COUNT(*) FILTER (
                            WHERE NOT ({CANCELLED_ORDER_SQL})
                        ) AS total_orders,
                        COUNT(*) FILTER (
                            WHERE ({CANCELLED_ORDER_SQL})
                        ) AS cancelled_orders,
                        COUNT(*) FILTER (
                            WHERE NOT ({CANCELLED_ORDER_SQL})
                              AND o.created_at >= CURRENT_DATE
                              AND o.created_at < CURRENT_DATE + INTERVAL '1 day'
                        ) AS orders_today,
                        COUNT(*) FILTER (
                            WHERE ({CANCELLED_ORDER_SQL})
                              AND o.created_at >= CURRENT_DATE
                              AND o.created_at < CURRENT_DATE + INTERVAL '1 day'
                        ) AS cancelled_orders_today
                    FROM vgf_orders o
                    WHERE o.is_active IS TRUE AND o.is_deleted IS FALSE
                )
                SELECT * FROM customer_counts CROSS JOIN ratings CROSS JOIN order_counts
                """
            )
            row = cur.fetchone()

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "kpis": [
            {
                "key": "total_customers",
                "title": "Total Customers",
                "value": row["total"],
                "new_this_week": row["new_this_week"],
                "new_last_week": row["new_last_week"],
                "change_this_week": row["new_this_week"] - row["new_last_week"],
                "source_table": "vgf_customers",
            },
            {
                "key": "inactive_customers",
                "title": "Inactive Customers",
                "value": row["inactive"],
                "percentage": round(row["inactive"] * 100 / row["total"], 1)
                if row["total"]
                else 0,
                "source_table": "vgf_customers",
            },
            {
                "key": "new_this_week",
                "title": "New This Week",
                "value": row["new_this_week"],
                "last_week": row["new_last_week"],
                "change": row["new_this_week"] - row["new_last_week"],
                "source_table": "vgf_customers",
            },
            {
                "key": "average_customer_rating",
                "title": "Average Customer Rating",
                "value": row["average_rating"],
                "review_count": row["review_count"],
                "source_table": "vgf_customer_feedback",
            },
            {
                "key": "total_orders",
                "title": "Total Orders",
                "value": row["total_orders"],
                "today": row["orders_today"],
                "source_table": "vgf_orders",
            },
            {
                "key": "cancelled_orders",
                "title": "Cancelled Orders",
                "value": row["cancelled_orders"],
                "today": row["cancelled_orders_today"],
                "source_table": "vgf_orders",
            },
        ],
    }


@app.get("/admin/customers/kpi/{card_key}")
async def customer_kpi_detail(
    card_key: str,
    feedback_id: int | None = Query(default=None, ge=1),
    search: str | None = Query(default=None, max_length=200),
    status: Literal["all", "active", "inactive"] = "all",
    joined_from: date | None = None,
    joined_to: date | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=25, ge=1, le=100),
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    del current_admin
    customer_filters = {
        "total_customers": ["c.is_deleted IS FALSE"],
        "inactive_customers": ["c.is_deleted IS FALSE", "c.is_active IS NOT TRUE"],
        "new_this_week": [
            "c.is_deleted IS FALSE",
            "c.created_at >= date_trunc('week', NOW())",
        ],
    }

    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            if card_key in customer_filters:
                where_clause, parameters = _customer_list_filters(
                    search,
                    status,
                    joined_from,
                    joined_to,
                )
                for condition in customer_filters[card_key]:
                    if condition != "c.is_deleted IS FALSE":
                        where_clause += f" AND {condition}"
                cur.execute(
                    f"SELECT COUNT(*) AS total FROM vgf_customers c WHERE {where_clause}",
                    parameters,
                )
                total = cur.fetchone()["total"]
                cur.execute(
                    f"""
                    SELECT c.customer_id, c.name, c.mobile, c.email,
                           c.loyalty_points, c.is_active, c.created_at,
                           COALESCE(order_counts.order_count, 0) AS order_count,
                           COALESCE(order_counts.cancelled_order_count, 0) AS cancelled_order_count
                    FROM vgf_customers c
                    LEFT JOIN (
                        SELECT o.customer_id,
                               COUNT(*) FILTER (
                                   WHERE NOT ({CANCELLED_ORDER_SQL})
                               ) AS order_count,
                               COUNT(*) FILTER (
                                   WHERE ({CANCELLED_ORDER_SQL})
                               ) AS cancelled_order_count
                        FROM vgf_orders o
                        WHERE o.is_active IS TRUE AND o.is_deleted IS FALSE
                        GROUP BY o.customer_id
                    ) order_counts USING (customer_id)
                    WHERE {where_clause}
                    ORDER BY c.created_at DESC, c.customer_id DESC
                    LIMIT %s OFFSET %s
                    """,
                    [*parameters, page_size, (page - 1) * page_size],
                )
                rows = cur.fetchall()
                titles = {
                    "total_customers": "Total Customers",
                    "inactive_customers": "Inactive Customers",
                    "new_this_week": "New This Week",
                }
                return {
                    "key": card_key,
                    "title": titles[card_key],
                    "source_table": "vgf_customers",
                    "count": total,
                    "data": rows,
                    "page": page,
                    "page_size": page_size,
                    "total_pages": (total + page_size - 1) // page_size,
                    "search": search or "",
                    "status": status,
                    "joined_from": joined_from.isoformat() if joined_from else "",
                    "joined_to": joined_to.isoformat() if joined_to else "",
                }

            if card_key == "average_customer_rating":
                cur.execute(
                    """
                    SELECT AVG(rating)::float AS average_rating, COUNT(*) AS review_count
                    FROM vgf_customer_feedback
                    WHERE is_active IS TRUE AND is_deleted IS FALSE
                    """
                )
                summary = cur.fetchone()
                conditions = ["f.is_active IS TRUE", "f.is_deleted IS FALSE"]
                feedback_parameters: list[Any] = []
                if search and search.strip():
                    pattern = f"%{search.strip()}%"
                    conditions.append(
                        "(c.name ILIKE %s OR f.comments ILIKE %s "
                        "OR CAST(f.order_id AS TEXT) ILIKE %s)"
                    )
                    feedback_parameters.extend((pattern, pattern, pattern))
                feedback_where = " AND ".join(conditions)
                cur.execute(
                    f"SELECT COUNT(*) AS total FROM vgf_customer_feedback f "
                    f"LEFT JOIN vgf_customers c ON c.customer_id = f.customer_id "
                    f"WHERE {feedback_where}",
                    feedback_parameters,
                )
                total = cur.fetchone()["total"]
                feedback_query = f"""
                    SELECT f.feedback_id, f.customer_id, c.name AS customer_name,
                           f.order_id, f.rating, f.comments, f.created_at
                    FROM vgf_customer_feedback f
                    LEFT JOIN vgf_customers c ON c.customer_id = f.customer_id
                    WHERE {feedback_where}
                """
                query_parameters = feedback_parameters.copy()
                if feedback_id is not None:
                    feedback_query += " ORDER BY (f.feedback_id = %s) DESC, "
                    query_parameters.append(feedback_id)
                else:
                    feedback_query += " ORDER BY "
                feedback_query += (
                    "f.created_at DESC, f.feedback_id DESC LIMIT %s OFFSET %s"
                )
                query_parameters.extend((page_size, (page - 1) * page_size))
                cur.execute(feedback_query, query_parameters)
                return {
                    "key": card_key,
                    "title": "Average Customer Rating",
                    "source_table": "vgf_customer_feedback",
                    "average_rating": summary["average_rating"],
                    "count": total,
                    "overall_review_count": summary["review_count"],
                    "data": cur.fetchall(),
                    "page": page,
                    "page_size": page_size,
                    "total_pages": (total + page_size - 1) // page_size,
                    "search": search or "",
                    "feedback_id": feedback_id,
                }

    raise HTTPException(status_code=404, detail="Customer KPI card not found")


@app.get("/admin/customers/orders")
async def list_customer_orders(
    order_type: Literal["orders", "cancelled"] = "orders",
    search: str | None = Query(default=None, max_length=200),
    customer_id: int | None = Query(default=None, ge=1),
    page: int = Query(default=1, ge=1),
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    del current_admin
    conditions = [
        "o.is_active IS TRUE",
        "o.is_deleted IS FALSE",
        (
            f"({CANCELLED_ORDER_SQL})"
            if order_type == "cancelled"
            else f"NOT ({CANCELLED_ORDER_SQL})"
        ),
    ]
    parameters: list[Any] = []
    if customer_id is not None:
        conditions.append("o.customer_id = %s")
        parameters.append(customer_id)
    if search and search.strip():
        pattern = f"%{search.strip()}%"
        conditions.append(
            "(c.name ILIKE %s OR CAST(o.order_id AS TEXT) ILIKE %s)"
        )
        parameters.extend((pattern, pattern))
    where_clause = " AND ".join(conditions)
    page_size = 25

    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                f"""
                SELECT COUNT(*) AS total
                FROM vgf_orders o
                LEFT JOIN vgf_customers c ON c.customer_id = o.customer_id
                WHERE {where_clause}
                """,
                parameters,
            )
            total = cur.fetchone()["total"]
            cur.execute(
                f"""
                SELECT o.order_id, o.customer_id, c.name AS customer_name,
                       o.order_status, o.cancellation_status,
                       o.total_amount AS amount, o.created_at AS date
                FROM vgf_orders o
                LEFT JOIN vgf_customers c ON c.customer_id = o.customer_id
                WHERE {where_clause}
                ORDER BY o.created_at DESC, o.order_id DESC
                LIMIT %s OFFSET %s
                """,
                [*parameters, page_size, (page - 1) * page_size],
            )
            orders = cur.fetchall()

    return {
        "order_type": order_type,
        "customer_id": customer_id,
        "items": orders,
        "page": page,
        "page_size": page_size,
        "total": total,
        "total_pages": (total + page_size - 1) // page_size,
    }


@app.get("/admin/quick-search")
async def quick_search(
    q: str = Query(min_length=2, max_length=200),
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    del current_admin
    pattern = f"%{q.strip()}%"
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                """
                SELECT customer_id, name, mobile, email
                FROM vgf_customers
                WHERE is_deleted IS FALSE
                  AND (name ILIKE %s OR mobile ILIKE %s OR email ILIKE %s)
                ORDER BY name, customer_id
                LIMIT 8
                """,
                (pattern, pattern, pattern),
            )
            customers = cur.fetchall()
            cur.execute(
                f"""
                SELECT o.order_id, o.customer_id, c.name AS customer_name,
                       o.order_status, o.total_amount AS amount, o.created_at AS date
                FROM vgf_orders o
                LEFT JOIN vgf_customers c ON c.customer_id = o.customer_id
                WHERE o.is_active IS TRUE AND o.is_deleted IS FALSE
                  AND NOT ({CANCELLED_ORDER_SQL})
                  AND CAST(o.order_id AS TEXT) ILIKE %s
                ORDER BY o.created_at DESC, o.order_id DESC
                LIMIT 8
                """,
                (pattern,),
            )
            orders = cur.fetchall()
    return {"customers": customers, "orders": orders}


@app.get("/admin/customers/{customer_id}")
async def get_customer(
    customer_id: int,
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    del current_admin
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            customer = _fetch_customer(cur, customer_id)
            if not customer:
                raise HTTPException(status_code=404, detail="Customer not found")

            cur.execute(
                """
                SELECT address_id, address_text, latitude, longitude, is_active
                FROM vgf_customer_addresses
                WHERE customer_id = %s AND is_deleted IS FALSE
                ORDER BY address_id
                """,
                (customer_id,),
            )
            addresses = cur.fetchall()
            cur.execute(
                """
                SELECT order_id, order_status, subtotal, gst_amount, delivery_charge,
                       discount_amount, total_amount, created_at, delivery_eta
                FROM vgf_orders
                WHERE customer_id = %s AND is_active IS TRUE AND is_deleted IS FALSE
                ORDER BY created_at DESC, order_id DESC
                """,
                (customer_id,),
            )
            orders = cur.fetchall()
            cur.execute(
                f"""
                SELECT COUNT(*) FILTER (
                           WHERE NOT ({CANCELLED_ORDER_SQL})
                       ) AS order_count,
                       COUNT(*) FILTER (
                           WHERE ({CANCELLED_ORDER_SQL})
                       ) AS cancelled_order_count
                FROM vgf_orders o
                WHERE o.customer_id = %s
                  AND o.is_active IS TRUE AND o.is_deleted IS FALSE
                """,
                (customer_id,),
            )
            order_counts = cur.fetchone()
            cur.execute(
                """
                SELECT AVG(rating)::float AS average_rating, COUNT(*) AS review_count
                FROM vgf_customer_feedback
                WHERE customer_id = %s AND is_active IS TRUE AND is_deleted IS FALSE
                """,
                (customer_id,),
            )
            rating = cur.fetchone()
            cur.execute(
                """
                SELECT a.id, a.action, a.old_data, a.new_data, a.changed_by,
                       a.added_by, a.created_at, u.full_name AS changed_by_name
                FROM vgf_customer_audit a
                LEFT JOIN vgf_users u ON u.user_id = a.changed_by
                WHERE a.customer_id = %s AND a.is_deleted IS FALSE
                ORDER BY a.created_at DESC
                """,
                (customer_id,),
            )
            audit_trail = cur.fetchall()

    return {
        "customer": customer,
        "addresses": addresses,
        "orders": orders,
        "metrics": {
            "order_count": order_counts["order_count"],
            "cancelled_order_count": order_counts["cancelled_order_count"],
            "loyalty_points": customer["loyalty_points"],
            "average_rating": rating["average_rating"],
            "review_count": rating["review_count"],
        },
        "audit_trail": audit_trail,
    }


@app.get("/admin/customers/{customer_id}/history/{history_key}")
async def get_customer_metric_history(
    customer_id: int,
    history_key: str,
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    del current_admin
    if history_key not in {
        "orders",
        "cancelled_orders",
        "loyalty_points",
        "average_rating",
        "reviews",
    }:
        raise HTTPException(status_code=404, detail="Customer history not found")

    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            customer = _fetch_customer(cur, customer_id)
            if not customer:
                raise HTTPException(status_code=404, detail="Customer not found")

            if history_key in {"orders", "cancelled_orders", "loyalty_points"}:
                cancelled_filter = (
                    f"AND ({CANCELLED_ORDER_SQL})"
                    if history_key == "cancelled_orders"
                    else f"AND NOT ({CANCELLED_ORDER_SQL})"
                )
                cur.execute(
                    f"""
                    SELECT o.order_id, o.created_at AS date,
                           o.order_status AS status, o.total_amount AS total
                    FROM vgf_orders o
                    WHERE o.customer_id = %s
                      AND o.is_active IS TRUE
                      AND o.is_deleted IS FALSE
                      {cancelled_filter}
                    ORDER BY o.created_at DESC, o.order_id DESC
                    """,
                    (customer_id,),
                )
                title = (
                    "Customer Cancelled Orders"
                    if history_key == "cancelled_orders"
                    else "Customer Orders"
                    if history_key == "orders"
                    else "Orders Related to Loyalty Points"
                )
                source_table = "vgf_orders"
            else:
                cur.execute(
                    """
                    SELECT feedback_id, order_id, created_at AS date,
                           rating, comments
                    FROM vgf_customer_feedback
                    WHERE customer_id = %s
                      AND is_active IS TRUE
                      AND is_deleted IS FALSE
                    ORDER BY created_at DESC, feedback_id DESC
                    """,
                    (customer_id,),
                )
                title = (
                    "Customer Average Rating"
                    if history_key == "average_rating"
                    else "Customer Reviews"
                )
                source_table = "vgf_customer_feedback"
            records = cur.fetchall()

    return {
        "key": history_key,
        "title": title,
        "source_table": source_table,
        "count": len(records),
        "data": records,
    }


@app.get("/admin/notifications")
async def list_admin_notifications(
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    admin_id = current_admin.get("admin_id")
    if not isinstance(admin_id, int):
        try:
            platform_admin_id = UUID(str(admin_id))
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=401, detail="Token missing admin ID") from exc

        with get_db_connection() as conn:
            with conn.cursor(row_factory=dict_row) as cur:
                cur.execute(
                    """
                    SELECT id, notification_type, title, message, created_at,
                           target_module, customer_id, order_id
                    FROM (
                        SELECT
                            'customer-' || c.customer_id::text AS id,
                            'new_customer'::text AS notification_type,
                            'New customer joined'::text AS title,
                            c.name || ' joined Vigo Feast.' AS message,
                            c.created_at,
                            'customers'::text AS target_module,
                            c.customer_id,
                            NULL::bigint AS order_id
                        FROM vgf_customers c
                        WHERE c.created_at >= NOW() - INTERVAL '7 days'
                          AND c.is_deleted IS NOT TRUE

                        UNION ALL

                        SELECT
                            'order-' || o.order_id::text,
                            'new_order'::text,
                            'New order placed'::text,
                            'Order #' || o.order_id::text || ' placed by ' || COALESCE(c.name, 'a customer') || '.',
                            o.created_at,
                            'customers'::text,
                            o.customer_id,
                            o.order_id
                        FROM vgf_orders o
                        LEFT JOIN vgf_customers c ON c.customer_id = o.customer_id
                        WHERE o.created_at >= NOW() - INTERVAL '7 days'
                          AND o.is_deleted IS NOT TRUE
                          AND LOWER(COALESCE(o.order_status, '')) IN ('placed', 'ordered')

                        UNION ALL

                        SELECT
                            'cancelled-order-' || o.order_id::text,
                            'cancelled_order'::text,
                            CASE
                                WHEN LOWER(COALESCE(o.order_status, '')) = 'rejected'
                                    THEN 'Order rejected'
                                ELSE 'Order cancelled'
                            END,
                            'Order #' || o.order_id::text || ' for ' || COALESCE(c.name, 'a customer') || ' was ' ||
                                CASE
                                    WHEN LOWER(COALESCE(o.order_status, '')) = 'rejected'
                                        THEN 'rejected.'
                                    ELSE 'cancelled.'
                                END,
                            o.created_at,
                            'customers'::text,
                            o.customer_id,
                            o.order_id
                        FROM vgf_orders o
                        LEFT JOIN vgf_customers c ON c.customer_id = o.customer_id
                        WHERE o.created_at >= NOW() - INTERVAL '7 days'
                          AND o.is_deleted IS NOT TRUE
                          AND (
                              LOWER(COALESCE(o.order_status, '')) IN ('cancelled', 'canceled', 'rejected')
                              OR LOWER(COALESCE(o.cancellation_status, '')) IN ('cancelled', 'canceled', 'rejected')
                          )

                        UNION ALL

                        SELECT
                            'rating-' || f.feedback_id::text,
                            'low_rating'::text,
                            'Low customer rating'::text,
                            COALESCE(c.name, 'A customer') || ' rated order #' || f.order_id::text ||
                                ' ' || f.rating::text || '/5.',
                            f.created_at,
                            'customers'::text,
                            f.customer_id,
                            f.order_id
                        FROM vgf_customer_feedback f
                        LEFT JOIN vgf_customers c ON c.customer_id = f.customer_id
                        WHERE f.created_at >= NOW() - INTERVAL '7 days'
                          AND f.is_deleted IS NOT TRUE
                          AND f.rating <= 3
                    ) AS recent_notifications
                    ORDER BY created_at DESC, id DESC
                    LIMIT 10
                    """
                )
                rows = cur.fetchall()

        items = [
            {
                "id": row["id"],
                "title": row["title"],
                "message": row["message"],
                "notification_type": row["notification_type"],
                "is_read": False,
                "created_at": row["created_at"],
                "target": {
                    "module": row["target_module"],
                    "customer_id": row["customer_id"],
                    "order_id": row["order_id"],
                },
            }
            for row in rows
        ]
        return {
            "items": items,
            "unread_count": len(items),
            "local_read_key": f"vigo_admin_read_notifications:{platform_admin_id}",
        }

    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                """
                SELECT notification_id, title, message, notification_type,
                       is_read, created_at, read_at
                FROM vgf_adm_notification
                WHERE admin_id = %s
                ORDER BY created_at DESC, notification_id DESC
                LIMIT 100
                """,
                (admin_id,),
            )
            notifications = cur.fetchall()
            cur.execute(
                """
                SELECT COUNT(*) AS unread_count
                FROM vgf_adm_notification
                WHERE admin_id = %s AND is_read IS FALSE
                """,
                (admin_id,),
            )
            unread_count = cur.fetchone()["unread_count"]

    return {"items": notifications, "unread_count": unread_count}


@app.post("/admin/notifications/{notification_id}/read")
async def mark_admin_notification_read(
    notification_id: str,
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    admin_id = current_admin.get("admin_id")
    if not isinstance(admin_id, int):
        try:
            UUID(str(admin_id))
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=401, detail="Token missing admin ID") from exc
        prefix, separator, raw_id = notification_id.rpartition("-")
        if not separator or prefix not in {"customer", "order", "cancelled-order", "rating"}:
            raise HTTPException(status_code=404, detail="Notification not found")
        try:
            if int(raw_id) < 1:
                raise ValueError
        except ValueError as exc:
            raise HTTPException(status_code=404, detail="Notification not found") from exc
        return {
            "notification": {"id": notification_id, "is_read": True},
            "local_only": True,
        }

    try:
        notification_id = int(notification_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail="Notification not found") from exc

    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                """
                UPDATE vgf_adm_notification
                SET is_read = TRUE, read_at = COALESCE(read_at, NOW())
                WHERE notification_id = %s AND admin_id = %s
                RETURNING notification_id, is_read, read_at
                """,
                (notification_id, admin_id),
            )
            notification = cur.fetchone()
            if not notification:
                raise HTTPException(status_code=404, detail="Notification not found")
            cur.execute(
                """
                SELECT COUNT(*) AS unread_count
                FROM vgf_adm_notification
                WHERE admin_id = %s AND is_read IS FALSE
                """,
                (admin_id,),
            )
            unread_count = cur.fetchone()["unread_count"]

    return {"notification": notification, "unread_count": unread_count}


@app.put("/admin/customers/{customer_id}")
async def update_customer(
    customer_id: int,
    payload: CustomerUpdateRequest,
    request: Request,
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    changes = payload.model_dump(exclude_unset=True)
    if not changes:
        raise HTTPException(status_code=422, detail="At least one field is required")
    if "name" in changes and (changes["name"] is None or not changes["name"].strip()):
        raise HTTPException(status_code=422, detail="name cannot be empty")

    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            customer = _fetch_customer(cur, customer_id, for_update=True)
            if not customer:
                raise HTTPException(status_code=404, detail="Customer not found")
            if "default_address_id" in changes and changes["default_address_id"] is not None:
                cur.execute(
                    """
                    SELECT 1 FROM vgf_customer_addresses
                    WHERE address_id = %s AND customer_id = %s AND is_deleted IS FALSE
                    """,
                    (changes["default_address_id"], customer_id),
                )
                if not cur.fetchone():
                    raise HTTPException(
                        status_code=422,
                        detail="default_address_id must belong to this customer",
                    )

            old_data = _customer_snapshot(customer)
            assignments = ", ".join(f"{column} = %s" for column in changes)
            values = [
                value.strip() if column == "name" and value is not None else value
                for column, value in changes.items()
            ]
            cur.execute(
                f"""
                UPDATE vgf_customers
                SET {assignments}, updated_at = NOW(),
                    last_updated_by = %s, last_updated_date = CURRENT_DATE,
                    last_updated_time = CURRENT_TIME, last_updated_ip = %s
                WHERE customer_id = %s
                RETURNING customer_id, user_id, name, mobile, email, default_address_id,
                          loyalty_points, created_at, updated_at, franchisee_id,
                          is_active, is_deleted
                """,
                [
                    *values,
                    str(current_admin.get("sub") or "admin"),
                    request.client.host if request.client else None,
                    customer_id,
                ],
            )
            updated = cur.fetchone()
            _write_customer_audit(
                cur,
                customer=updated,
                current_admin=current_admin,
                request=request,
                action="customer.updated",
                old_data=old_data,
                new_data=_customer_snapshot(updated),
            )
            _write_admin_log(
                cur,
                current_admin=current_admin,
                action="customer.updated",
                entity_type="customer",
                description=f"Edited customer {updated['name']}.",
                ip_address=request.client.host if request.client else None,
                entity_name=updated["name"],
                old_data=old_data,
                new_data=_customer_snapshot(updated),
            )
    return {"customer": updated}


@app.put("/admin/customers/{customer_id}/status")
async def update_customer_status(
    customer_id: int,
    payload: CustomerStatusRequest,
    request: Request,
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            customer = _fetch_customer(cur, customer_id, for_update=True)
            if not customer:
                raise HTTPException(status_code=404, detail="Customer not found")
            if customer["is_active"] is payload.is_active:
                return {"customer": customer}

            old_data = _customer_snapshot(customer)
            cur.execute(
                """
                UPDATE vgf_customers
                SET is_active = %s, updated_at = NOW(),
                    last_updated_by = %s, last_updated_date = CURRENT_DATE,
                    last_updated_time = CURRENT_TIME, last_updated_ip = %s
                WHERE customer_id = %s
                RETURNING customer_id, user_id, name, mobile, email, default_address_id,
                          loyalty_points, created_at, updated_at, franchisee_id,
                          is_active, is_deleted
                """,
                (
                    payload.is_active,
                    str(current_admin.get("sub") or "admin"),
                    request.client.host if request.client else None,
                    customer_id,
                ),
            )
            updated = cur.fetchone()
            _write_customer_audit(
                cur,
                customer=updated,
                current_admin=current_admin,
                request=request,
                action="customer.status_changed",
                old_data=old_data,
                new_data=_customer_snapshot(updated),
            )
            _write_admin_log(
                cur,
                current_admin=current_admin,
                action=(
                    "customer.activated"
                    if payload.is_active
                    else "customer.deactivated"
                ),
                entity_type="customer",
                description=(
                    f"{'Activated' if payload.is_active else 'Deactivated'} "
                    f"customer {updated['name']}."
                ),
                ip_address=request.client.host if request.client else None,
                entity_name=updated["name"],
                old_data=old_data,
                new_data=_customer_snapshot(updated),
            )
    return {"customer": updated}


@app.delete("/admin/customers/{customer_id}")
async def delete_customer(
    customer_id: int,
    request: Request,
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            customer = _fetch_customer(cur, customer_id, for_update=True)
            if not customer:
                raise HTTPException(status_code=404, detail="Customer not found")

            old_data = _customer_snapshot(customer)
            cur.execute(
                """
                UPDATE vgf_customers
                SET is_active = FALSE, is_deleted = TRUE, updated_at = NOW(),
                    last_updated_by = %s, last_updated_date = CURRENT_DATE,
                    last_updated_time = CURRENT_TIME, last_updated_ip = %s
                WHERE customer_id = %s
                RETURNING customer_id, user_id, name, mobile, email, default_address_id,
                          loyalty_points, created_at, updated_at, franchisee_id,
                          is_active, is_deleted
                """,
                (
                    str(current_admin.get("sub") or "admin"),
                    request.client.host if request.client else None,
                    customer_id,
                ),
            )
            deleted = cur.fetchone()
            _write_customer_audit(
                cur,
                customer=deleted,
                current_admin=current_admin,
                request=request,
                action="customer.deleted",
                old_data=old_data,
                new_data=_customer_snapshot(deleted),
            )
            _write_admin_log(
                cur,
                current_admin=current_admin,
                action="customer.deleted",
                entity_type="customer",
                description=f"Deleted customer {deleted['name']}.",
                ip_address=request.client.host if request.client else None,
                entity_name=deleted["name"],
                old_data=old_data,
                new_data=_customer_snapshot(deleted),
            )
    return {"message": "Customer deleted", "customer_id": customer_id}


@app.get("/health")
async def health() -> JSONResponse:
    try:
        await asyncio.to_thread(check_db_connection)
    except Exception as exc:  # pragma: no cover - runtime dependency check
        return JSONResponse(
            status_code=503,
            content={
                "status": "error",
                "database": "disconnected",
                "detail": str(exc),
            },
        )

    return JSONResponse(
        content={
            "status": "ok",
            "service": "Vigo Feast Admin Console",
            "database": "connected",
        }
    )


@app.post("/admin/login")
async def admin_login(payload: AdminPasswordLoginRequest, request: Request):
    admin = find_admin_by_username(payload.username)
    if not admin:
        raise HTTPException(status_code=401, detail="Invalid username or password")
    if admin.get("is_active") is False:
        raise HTTPException(status_code=403, detail="Admin account is inactive")
    if not verify_password(payload.password, admin.get("password_hash")):
        raise HTTPException(status_code=401, detail="Invalid username or password")

    ip_address = request.client.host if request.client else "127.0.0.1"
    user_agent = request.headers.get("user-agent", "unknown")
    session = create_admin_session(admin, ip_address, user_agent)
    token = create_admin_token(admin, session["session_id"])
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            _write_admin_log(
                cur,
                current_admin={
                    "session_id": session["session_id"],
                    "sub": admin.get("username") or admin.get("email") or "admin",
                },
                action="admin.login",
                entity_type="admin_session",
                description="Admin signed in.",
                ip_address=ip_address,
                entity_id=UUID(session["session_id"]),
                entity_name=admin.get("full_name"),
            )
    return {
        "token": token,
        "admin": {
            "admin_id": admin.get("admin_id"),
            "full_name": admin.get("full_name"),
            "email": admin.get("email"),
            "mobile_no": admin.get("mobile_no"),
            "username": admin.get("username"),
            "role": admin.get("role"),
        },
        "session": session,
        "message": "Login successful",
    }


@app.post("/admin/login/otp")
async def admin_login_otp(payload: AdminOtpRequest):
    admin = find_admin_by_mobile(payload.mobile_number)
    if not admin:
        raise HTTPException(status_code=404, detail="No admin found for this mobile number")
    if admin.get("is_active") is False:
        raise HTTPException(status_code=403, detail="Admin account is inactive")

    otp_code = generate_otp()
    try:
        nimbus_message = send_otp_message(
            payload.mobile_number.strip(), otp_code, payload.language
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    saved = save_otp_code(payload.mobile_number, otp_code, admin)
    response = {
        "message": nimbus_message,
        "mobile_number": saved["mobile_number"],
        "expires_in_seconds": 600,
    }
    if settings.app_env == "local":
        response["otp_code"] = otp_code
    return response


@app.post("/admin/login/verify")
async def admin_login_verify(payload: AdminOtpVerifyRequest, request: Request):
    verified = verify_otp_code(payload.mobile_number, payload.otp_code)
    if not verified or not verified.get("admin"):
        raise HTTPException(status_code=401, detail="Invalid or expired OTP code")

    admin = verified["admin"]
    if admin.get("is_active") is False:
        raise HTTPException(status_code=403, detail="Admin account is inactive")

    ip_address = request.client.host if request.client else "127.0.0.1"
    user_agent = request.headers.get("user-agent", "unknown")
    session = create_admin_session(admin, ip_address, user_agent)
    token = create_admin_token(admin, session["session_id"])
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            _write_admin_log(
                cur,
                current_admin={
                    "session_id": session["session_id"],
                    "sub": admin.get("username") or admin.get("email") or "admin",
                },
                action="admin.login",
                entity_type="admin_session",
                description="Admin signed in.",
                ip_address=ip_address,
                entity_id=UUID(session["session_id"]),
                entity_name=admin.get("full_name"),
            )
    return {
        "token": token,
        "admin": {
            "admin_id": admin.get("admin_id"),
            "full_name": admin.get("full_name"),
            "email": admin.get("email"),
            "mobile_no": admin.get("mobile_no"),
            "username": admin.get("username"),
            "role": admin.get("role"),
        },
        "session": session,
        "message": "OTP verified and token issued",
    }


@app.get("/admin/dashboard")
async def admin_dashboard(current_admin: dict[str, Any] = Depends(require_admin_auth)):
    del current_admin
    return _build_dashboard_payload()


@app.get("/admin/me")
async def admin_profile(current_admin: dict[str, Any] = Depends(require_admin_auth)):
    admin_id = current_admin.get("admin_id")
    subject = current_admin.get("sub")
    with get_db_connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            admin = None
            platform_admin_id = None
            if admin_id is not None:
                try:
                    platform_admin_id = UUID(str(admin_id))
                except (TypeError, ValueError):
                    pass
                if platform_admin_id is not None:
                    cur.execute(
                        "SELECT name AS full_name, phone AS mobile_no, role FROM vgf_platform_admins WHERE id = %s",
                        (platform_admin_id,),
                    )
                    admin = cur.fetchone()
                else:
                    cur.execute(
                        "SELECT full_name, mobile_no, role FROM vgf_admin WHERE admin_id = %s",
                        (admin_id,),
                    )
                    admin = cur.fetchone()
            if admin is None and platform_admin_id is None and subject:
                cur.execute(
                    """
                    SELECT full_name, mobile_no, role
                    FROM vgf_admin
                    WHERE username = %s OR email = %s OR mobile_no = %s
                    LIMIT 1
                    """,
                    (subject, subject, subject),
                )
                admin = cur.fetchone()
    if admin is None:
        raise HTTPException(status_code=404, detail="Admin profile not found")
    return {
        "full_name": admin["full_name"],
        "mobile_no": admin["mobile_no"],
        "role": admin["role"] or current_admin.get("role"),
    }


@app.post("/admin/logout")
async def admin_logout(current_admin: dict[str, Any] = Depends(require_admin_auth)):
    session_id = current_admin.get("session_id")
    if not isinstance(session_id, str):
        raise HTTPException(status_code=401, detail="Token missing session ID")
    try:
        session_id = UUID(session_id)
    except ValueError as exc:
        raise HTTPException(status_code=401, detail="Token has an invalid session ID") from exc

    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE vgf_adm_sessions
                SET is_active = FALSE
                WHERE id = %s AND is_deleted IS FALSE
                """,
                (session_id,),
            )
            if cur.rowcount == 0:
                raise HTTPException(status_code=404, detail="Session not found")
            _write_admin_log(
                cur,
                current_admin=current_admin,
                action="admin.logout",
                entity_type="admin_session",
                description="Admin signed out.",
                ip_address=None,
                entity_id=session_id,
            )
    return {"message": "Logged out"}


@app.get("/admin/dashboard/kpi/{card_key}")
async def admin_dashboard_kpi(
    card_key: str,
    order_id: int | None = Query(default=None, ge=1),
    current_admin: dict[str, Any] = Depends(require_admin_auth),
):
    del current_admin
    if card_key == "live_orders" and order_id is not None:
        return _kpi_detail_live_orders(order_id)
    return _get_kpi_detail(card_key)
